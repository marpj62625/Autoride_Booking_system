import sys, os, datetime, io
sys.path.append(os.path.abspath('backend'))
from database import get_connection
from psycopg.rows import dict_row

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

conn = get_connection()
cur = conn.cursor(row_factory=dict_row)

date_from = '2026-01-01'
date_to = '2026-12-31'
filter_period = 'year'
is_year_view = True

# 1. Summary stats
rev_query = """
    SELECT 
        SUM(CASE WHEN b.status = 'Cancelled' THEN (COALESCE(b.amount_paid, 0) - COALESCE(b.refund_amount, 0)) ELSE b.total_price END) as total_revenue,
        COUNT(b.id) as total_bookings
    FROM bookings b 
    JOIN vehicles v ON b.vehicle_id = v.id 
    WHERE (b.status != 'Cancelled' OR (b.status = 'Cancelled' AND COALESCE(b.amount_paid, 0) > COALESCE(b.refund_amount, 0)))
      AND b.start_date >= %s AND b.start_date <= %s
"""
cur.execute(rev_query, (date_from, date_to))
basic_stats = cur.fetchone()
total_revenue = float(basic_stats['total_revenue'] or 0)
total_bookings = int(basic_stats['total_bookings'] or 0)

# Expenses
exp_query = """
    SELECT COALESCE(SUM(amount), 0) as total_expenses 
    FROM vehicle_expenses 
    WHERE expense_date >= %s AND expense_date <= %s
"""
cur.execute(exp_query, (date_from, date_to))
total_expenses = float(cur.fetchone()['total_expenses'] or 0)
net_profit = total_revenue - total_expenses
avg_rev = math_floor = int(total_revenue / total_bookings) if total_bookings > 0 else 0

# Active fleet
cur.execute("SELECT COUNT(*) as count FROM vehicles WHERE status != 'Inactive'")
active_fleet = int(cur.fetchone()['count'] or 0)

# Trend
date_format = 'YYYY-MM' if is_year_view else 'YYYY-MM-DD'
trend_query = f"""
    SELECT 
        TO_CHAR(b.start_date, '{date_format}') as day,
        SUM(CASE WHEN b.status = 'Cancelled' THEN (COALESCE(b.amount_paid, 0) - COALESCE(b.refund_amount, 0)) ELSE b.total_price END) as amount,
        COUNT(b.id) as booking_count
    FROM bookings b
    JOIN vehicles v ON b.vehicle_id = v.id
    WHERE (b.status != 'Cancelled' OR (b.status = 'Cancelled' AND COALESCE(b.amount_paid, 0) > COALESCE(b.refund_amount, 0)))
      AND b.start_date >= %s AND b.start_date <= %s
    GROUP BY day ORDER BY day ASC
"""
cur.execute(trend_query, (date_from, date_to))
raw_trend = cur.fetchall()

target_year = date_from[:4] if date_from else str(datetime.date.today().year)
month_names = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
trend_map = {r['day']: r for r in raw_trend}
monthly_trend = []
for idx, m_num in enumerate(['01', '02', '03', '04', '05', '06', '07', '08', '09', '10', '11', '12']):
    m_key = f"{target_year}-{m_num}"
    if m_key in trend_map:
        row = trend_map[m_key]
        monthly_trend.append((month_names[idx], int(row.get('booking_count') or 0), float(row.get('amount') or 0)))
    else:
        monthly_trend.append((month_names[idx], 0, 0.0))

# Top vehicles
top_query = """
    SELECT v.brand, v.model, v.plate_number, COUNT(b.id) as booking_count, 
           COALESCE(SUM(CASE WHEN b.status = 'Cancelled' THEN (COALESCE(b.amount_paid, 0) - COALESCE(b.refund_amount, 0)) ELSE b.total_price END), 0) as revenue
    FROM vehicles v
    JOIN bookings b ON v.id = b.vehicle_id
    WHERE (b.status != 'Cancelled' OR (b.status = 'Cancelled' AND COALESCE(b.amount_paid, 0) > COALESCE(b.refund_amount, 0)))
      AND b.start_date >= %s AND b.start_date <= %s
    GROUP BY v.id, v.brand, v.model, v.plate_number ORDER BY revenue DESC LIMIT 5
"""
cur.execute(top_query, (date_from, date_to))
top_vehs = cur.fetchall()

# Full bookings list
bk_query = """
    SELECT b.id, b.customer_name, b.customer_phone, v.brand, v.model, v.plate_number, 
           b.start_date, b.end_date, b.total_price, b.status, b.payment_status
    FROM bookings b
    JOIN vehicles v ON b.vehicle_id = v.id
    WHERE (b.status != 'Cancelled' OR (b.status = 'Cancelled' AND COALESCE(b.amount_paid, 0) > COALESCE(b.refund_amount, 0)))
      AND b.start_date >= %s AND b.start_date <= %s
    ORDER BY b.start_date ASC, b.id ASC
"""
cur.execute(bk_query, (date_from, date_to))
bookings_list = cur.fetchall()

# Expenses list
exp_list_query = """
    SELECT e.expense_date, v.brand, v.model, v.plate_number, e.expense_type, e.amount, e.description, u.full_name as recorder_name
    FROM vehicle_expenses e
    JOIN vehicles v ON e.vehicle_id = v.id
    LEFT JOIN users u ON e.recorded_by = u.id
    WHERE e.expense_date >= %s AND e.expense_date <= %s
    ORDER BY e.expense_date DESC
"""
cur.execute(exp_list_query, (date_from, date_to))
expenses_list = cur.fetchall()

conn.close()

# BUILD WORKBOOK
wb = openpyxl.Workbook()

# Sheet 1: Executive Dashboard
ws_dash = wb.active
ws_dash.title = "Executive Summary"
ws_dash.views.sheetView[0].showGridLines = True

# Title banner
ws_dash["A1"] = "AUTORIDE SYSTEM — EXECUTIVE PERFORMANCE REPORT"
ws_dash["A1"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
ws_dash["A1"].fill = PatternFill(start_color="00B14F", end_color="00B14F", fill_type="solid")
ws_dash["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_dash.merge_cells("A1:K2")

ws_dash["A3"] = f"Period: {date_from} to {date_to} | Generated: {datetime.datetime.now().strftime('%b %d, %Y %I:%M %p')}"
ws_dash["A3"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

# Summary KPI Table
kpis = [
    ("TOTAL REVENUE", f"PHP {total_revenue:,.2f}"),
    ("TOTAL BOOKINGS", str(total_bookings)),
    ("AVG REVENUE / BOOKING", f"PHP {avg_rev:,.2f}"),
    ("TOTAL EXPENSES", f"PHP {total_expenses:,.2f}"),
    ("NET PROFIT", f"PHP {net_profit:,.2f}"),
    ("ACTIVE FLEET", str(active_fleet))
]

ws_dash["A5"] = "KPI METRIC"
ws_dash["B5"] = "VALUE"
ws_dash["A5"].font = Font(bold=True, color="FFFFFF")
ws_dash["B5"].font = Font(bold=True, color="FFFFFF")
ws_dash["A5"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
ws_dash["B5"].fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

thin_border = Border(left=Side(style='thin', color='CBD5E1'), right=Side(style='thin', color='CBD5E1'),
                     top=Side(style='thin', color='CBD5E1'), bottom=Side(style='thin', color='CBD5E1'))

for i, (k, v) in enumerate(kpis, start=6):
    ws_dash[f"A{i}"] = k
    ws_dash[f"B{i}"] = v
    ws_dash[f"A{i}"].font = Font(bold=True, size=10)
    ws_dash[f"B{i}"].font = Font(bold=True, size=10, color="00B14F" if ("PROFIT" in k or "REVENUE" in k) else "000000")
    ws_dash[f"A{i}"].border = thin_border
    ws_dash[f"B{i}"].border = thin_border

# Sheet 2: Data Sources for Charts
ws_data = wb.create_sheet(title="Chart_Data")
ws_data.views.sheetView[0].showGridLines = True

# Add Monthly Trend Data
ws_data.append(["Month", "Bookings", "Revenue (PHP)"])
for m_row in monthly_trend:
    ws_data.append(list(m_row))

for col in range(1, 4):
    cell = ws_data.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

# Add Top Vehicles Data in Chart_Data
ws_data.cell(row=1, column=5, value="Vehicle")
ws_data.cell(row=1, column=6, value="Revenue (PHP)")
ws_data.cell(row=1, column=5).font = Font(bold=True, color="FFFFFF")
ws_data.cell(row=1, column=6).font = Font(bold=True, color="FFFFFF")
ws_data.cell(row=1, column=5).fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
ws_data.cell(row=1, column=6).fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

for r_idx, tv in enumerate(top_vehs, start=2):
    ws_data.cell(row=r_idx, column=5, value=f"{tv['brand']} {tv['model']} ({tv['plate_number']})")
    ws_data.cell(row=r_idx, column=6, value=float(tv['revenue'] or 0))

# 1. Line Chart: Revenue Trend
line_chart = LineChart()
line_chart.title = "Revenue Trend (2026)"
line_chart.style = 13
line_chart.y_axis.title = "Revenue (PHP)"
line_chart.x_axis.title = "Month"
line_chart.width = 17
line_chart.height = 10.5

rev_data_ref = Reference(ws_data, min_col=3, min_row=1, max_row=13)
cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=13)
line_chart.add_data(rev_data_ref, titles_from_data=True)
line_chart.set_categories(cats_ref)
ws_dash.add_chart(line_chart, "D5")

# 2. Bar Chart: Bookings Detail
bar_chart = BarChart()
bar_chart.type = "col"
bar_chart.style = 10
bar_chart.title = "Bookings Detail (2026)"
bar_chart.y_axis.title = "Number of Bookings"
bar_chart.x_axis.title = "Month"
bar_chart.width = 17
bar_chart.height = 10.5

b_data_ref = Reference(ws_data, min_col=2, min_row=1, max_row=13)
bar_chart.add_data(b_data_ref, titles_from_data=True)
bar_chart.set_categories(cats_ref)
ws_dash.add_chart(bar_chart, "D23")

# 3. Pie Chart: Top Vehicles
pie_chart = PieChart()
pie_chart.title = "Top 5 Vehicles by Revenue"
pie_chart.width = 14
pie_chart.height = 10.5

pie_data = Reference(ws_data, min_col=6, min_row=1, max_row=len(top_vehs)+1)
pie_cats = Reference(ws_data, min_col=5, min_row=2, max_row=len(top_vehs)+1)
pie_chart.add_data(pie_data, titles_from_data=True)
pie_chart.set_categories(pie_cats)
ws_dash.add_chart(pie_chart, "A14")

# Sheet 3: All Bookings List
ws_bks = wb.create_sheet(title="All Bookings")
ws_bks.views.sheetView[0].showGridLines = True
ws_bks.append(["Booking ID", "Customer Name", "Phone", "Vehicle", "Plate Number", "Start Date", "End Date", "Total Price (PHP)", "Status", "Payment Status"])
for col in range(1, 11):
    cell = ws_bks.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

for b in bookings_list:
    ws_bks.append([
        b['id'], b['customer_name'], b['customer_phone'] or 'N/A',
        f"{b['brand']} {b['model']}", b['plate_number'],
        str(b['start_date']), str(b['end_date']),
        float(b['total_price'] or 0), b['status'], b['payment_status']
    ])

# Sheet 4: Expenses Log
ws_exp = wb.create_sheet(title="Expenses & Maintenance")
ws_exp.views.sheetView[0].showGridLines = True
ws_exp.append(["Date", "Vehicle", "Plate Number", "Expense Type", "Amount (PHP)", "Description", "Recorded By"])
for col in range(1, 8):
    cell = ws_exp.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

for e in expenses_list:
    ws_exp.append([
        str(e['expense_date']), f"{e['brand']} {e['model']}", e['plate_number'] or 'N/A',
        e['expense_type'], float(e['amount'] or 0), e['description'] or '', e['recorder_name'] or 'N/A'
    ])

# Auto-fit columns across all sheets
for ws in [ws_dash, ws_data, ws_bks, ws_exp]:
    for col in ws.columns:
        max_len = 0
        for cell in col:
            v_str = str(cell.value or '')
            if len(v_str) > max_len and len(v_str) < 50:
                max_len = len(v_str)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

wb.save("scratch/test_full_executive_report.xlsx")
print("SUCCESS: Full executive report generated at scratch/test_full_executive_report.xlsx")
