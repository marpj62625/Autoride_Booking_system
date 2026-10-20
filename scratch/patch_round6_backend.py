import sys

BACKEND_PATH = "backend/app.py"

with open(BACKEND_PATH, "r", encoding="utf-8") as f:
    content = f.read()

orig_len = len(content)

# Target 1: serve_uploads (~Line 1970)
p1_old = """@app.route('/uploads/<path:filename>')
@app.route('/api/uploads/<path:filename>')
def serve_uploads(filename):

    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)"""

p1_new = """@app.route('/uploads/<path:filename>')
@app.route('/api/uploads/<path:filename>')
def serve_uploads(filename):
    if '..' in filename or filename.startswith(('/', '\\\\')):
        return jsonify({"error": "Invalid file path"}), 400
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename)"""

assert p1_old in content, "Target 1 (serve_uploads) pattern not found"
content = content.replace(p1_old, p1_new, 1)

# Target 2: check_verify_status (~Line 3015)
p2_old = """def check_verify_status():

    user_id = request.args.get('user_id')

    if not user_id: return jsonify({"error": "Missing user_id"}), 400

    try:"""

p2_new = """def check_verify_status():

    user_id = request.args.get('user_id')

    if not user_id: return jsonify({"error": "Missing user_id"}), 400
    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p2_old in content, "Target 2 (check_verify_status) pattern not found"
content = content.replace(p2_old, p2_new, 1)

# Target 3: get_user_points (~Line 4094)
p3_old = """def get_user_points():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "User ID required"}), 400

    try:"""

p3_new = """def get_user_points():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "User ID required"}), 400
    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p3_old in content, "Target 3 (get_user_points) pattern not found"
content = content.replace(p3_old, p3_new, 1)

# Target 4: get_vehicle (~Line 5174)
p4_old = """def get_vehicle(vehicle_id):

    user_id = request.args.get('user_id')

    try:"""

p4_new = """def get_vehicle(vehicle_id):

    user_id = request.args.get('user_id')
    if user_id:
        try:
            user_id = int(user_id)
        except (ValueError, TypeError):
            user_id = None

    try:"""

assert p4_old in content, "Target 4 (get_vehicle) pattern not found"
content = content.replace(p4_old, p4_new, 1)

# Target 5: update_image_order (~Line 5397)
p5_old = """def update_image_order(vehicle_id):

    data = request.json

    order = data.get('order') # Expecting list of image IDs in order

    if not order:

        return jsonify({"error": "Order list is required"}), 400

    

    try:"""

p5_new = """def update_image_order(vehicle_id):
    is_authorized, err_resp = verify_admin_access()
    if not is_authorized:
        return err_resp

    data = request.json or {}

    order = data.get('order') # Expecting list of image IDs in order

    if not order:

        return jsonify({"error": "Order list is required"}), 400

    if not isinstance(order, list):
        return jsonify({"error": "Order must be a list of image IDs"}), 400
    if len(order) > 50:
        return jsonify({"error": "Order list exceeds maximum limit (50)"}), 400

    try:"""

assert p5_old in content, "Target 5 (update_image_order) pattern not found"
content = content.replace(p5_old, p5_new, 1)

# Target 6: profile (~Line 6202)
p6_old = """def profile():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id is required"}), 400

        

    try:"""

p6_new = """def profile():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id is required"}), 400

    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p6_old in content, "Target 6 (profile) pattern not found"
content = content.replace(p6_old, p6_new, 1)

# Target 7: get_favorites (~Line 6487)
p7_old = """def get_favorites():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id is required"}), 400

        

    try:"""

p7_new = """def get_favorites():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id is required"}), 400

    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p7_old in content, "Target 7 (get_favorites) pattern not found"
content = content.replace(p7_old, p7_new, 1)

# Target 8: get_saved_payments (~Line 6800)
p8_old = """def get_saved_payments():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id is required"}), 400

        

    try:"""

p8_new = """def get_saved_payments():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id is required"}), 400

    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p8_old in content, "Target 8 (get_saved_payments) pattern not found"
content = content.replace(p8_old, p8_new, 1)

# Target 9: driver_status (~Line 9659)
p9_old = """def driver_status():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id required"}), 400

        

    try:"""

p9_new = """def driver_status():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id required"}), 400

    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p9_old in content, "Target 9 (driver_status) pattern not found"
content = content.replace(p9_old, p9_new, 1)

# Target 10: get_driver_bookings (~Line 9697)
p10_old = """def get_driver_bookings():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id required"}), 400

        

    try:"""

p10_new = """def get_driver_bookings():

    user_id = request.args.get('user_id')

    if not user_id:

        return jsonify({"error": "user_id required"}), 400

    try:
        user_id = int(user_id)
    except (ValueError, TypeError):
        return jsonify({"error": "Invalid user_id"}), 400

    try:"""

assert p10_old in content, "Target 10 (get_driver_bookings) pattern not found"
content = content.replace(p10_old, p10_new, 1)

# Target 11: get_detailed_stats (~Line 9860)
p11_old = """    admin_id = request.args.get('admin_id')
    chart_type = request.args.get('type')
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    status_filter = request.args.get('status')
    vehicle_type = request.args.get('vehicle_type')
    filter_period = request.args.get('filter')"""

p11_new = """    admin_id = request.args.get('admin_id')
    try:
        admin_id = int(admin_id) if admin_id else None
    except (ValueError, TypeError):
        admin_id = None
    chart_type = str(request.args.get('type') or '').strip()[:50] or None
    date_from = str(request.args.get('date_from') or '').strip()[:10] or None
    date_to = str(request.args.get('date_to') or '').strip()[:10] or None
    status_filter = str(request.args.get('status') or '').strip()[:30] or None
    vehicle_type = str(request.args.get('vehicle_type') or '').strip()[:50] or None
    filter_period = str(request.args.get('filter') or '').strip()[:30] or None"""

assert p11_old in content, "Target 11 (get_detailed_stats) pattern not found"
content = content.replace(p11_old, p11_new, 1)

# Target 12: export_reports_excel (~Line 10035)
p12_old = """def export_reports_excel():
    \"\"\"Generates an executive Excel (.xlsx) report complete with native interactive charts.\"\"\"
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import io

    admin_id = request.args.get('admin_id')
    date_from = request.args.get('date_from')
    date_to = request.args.get('date_to')
    filter_period = request.args.get('filter', 'year')
    status_filter = request.args.get('status')
    vehicle_type = request.args.get('vehicle_type')"""

p12_new = """def export_reports_excel():
    \"\"\"Generates an executive Excel (.xlsx) report complete with native interactive charts.\"\"\"
    is_authorized, err_resp = verify_admin_access()
    if not is_authorized:
        return err_resp

    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import io

    admin_id = request.args.get('admin_id')
    try:
        admin_id = int(admin_id) if admin_id else None
    except (ValueError, TypeError):
        admin_id = None
    date_from = str(request.args.get('date_from') or '').strip()[:10] or None
    date_to = str(request.args.get('date_to') or '').strip()[:10] or None
    filter_period = str(request.args.get('filter', 'year') or 'year').strip()[:20]
    status_filter = str(request.args.get('status') or '').strip()[:30] or None
    vehicle_type = str(request.args.get('vehicle_type') or '').strip()[:50] or None"""

assert p12_old in content, "Target 12 (export_reports_excel) pattern not found"
content = content.replace(p12_old, p12_new, 1)

# Target 13: report_top_vehicles (~Line 10622)
p13_old = """    admin_id = request.args.get('admin_id')

    try:"""

p13_new = """    admin_id = request.args.get('admin_id')
    try:
        admin_id = int(admin_id) if admin_id else None
    except (ValueError, TypeError):
        admin_id = None

    try:"""

assert p13_old in content, "Target 13 (report_top_vehicles) pattern not found"
content = content.replace(p13_old, p13_new, 1)

# Target 14: export_maintenance_data (~Line 11748)
p14_old = """def export_maintenance_data(data_type):
    \"\"\"Admin endpoint to download JSON export backup of archived items or logs.\"\"\"
    import json
    from flask import Response
    from datetime import datetime

    try:"""

p14_new = """def export_maintenance_data(data_type):
    \"\"\"Admin endpoint to download JSON export backup of archived items or logs.\"\"\"
    is_authorized, err_resp = verify_admin_access()
    if not is_authorized:
        return err_resp

    if data_type not in ['archived_bookings', 'activity_logs', 'notifications']:
        return jsonify({"error": "Invalid export type. Options: archived_bookings, activity_logs, notifications"}), 400

    import json
    from flask import Response
    from datetime import datetime

    try:"""

assert p14_old in content, "Target 14 (export_maintenance_data) pattern not found"
content = content.replace(p14_old, p14_new, 1)

# Target 15: get_vehicle_details_v2 (~Line 13351)
p15_old = """def get_vehicle_details_v2(vehicle_id):


    user_id = request.args.get('user_id')

    try:"""

p15_new = """def get_vehicle_details_v2(vehicle_id):


    user_id = request.args.get('user_id')
    if user_id:
        try:
            user_id = int(user_id)
        except (ValueError, TypeError):
            user_id = None

    try:"""

assert p15_old in content, "Target 15 (get_vehicle_details_v2) pattern not found"
content = content.replace(p15_old, p15_new, 1)

# Target 16: bulk_update_settings (~Line 13813)
p16_old = """        updates = data.get('settings', [])

        if not updates:
            return jsonify({"error": "Missing settings payload"}), 400"""

p16_new = """        updates = data.get('settings', [])

        if not updates:
            return jsonify({"error": "Missing settings payload"}), 400

        if not isinstance(updates, list):
            return jsonify({"error": "Settings payload must be a list"}), 400
        if len(updates) > 100:
            return jsonify({"error": "Settings batch exceeds maximum limit (100)"}), 400"""

assert p16_old in content, "Target 16 (bulk_update_settings) pattern not found"
content = content.replace(p16_old, p16_new, 1)

# Target 17: confirm_fleet_calendar_import (~Line 17402)
p17_old = """def confirm_fleet_calendar_import():
    \"\"\"
    Imports the verified records into the database with smart renter linking.
    \"\"\"
    data = request.json or {}
    records = data.get('records') or []
    auto_create_renters = data.get('auto_create_renters', True)
    requester_id = data.get('requester_id', 1)

    if not records:
        return jsonify({'error': 'No records selected for import'}), 400"""

p17_new = """def confirm_fleet_calendar_import():
    \"\"\"
    Imports the verified records into the database with smart renter linking.
    \"\"\"
    is_authorized, err_resp = verify_admin_access()
    if not is_authorized:
        return err_resp

    data = request.json or {}
    records = data.get('records') or []
    auto_create_renters = data.get('auto_create_renters', True)
    requester_id = data.get('requester_id', 1)
    try:
        requester_id = int(requester_id)
    except (ValueError, TypeError):
        requester_id = 1

    if not records:
        return jsonify({'error': 'No records selected for import'}), 400

    if not isinstance(records, list):
        return jsonify({'error': 'Records must be a list'}), 400
    if len(records) > 2000:
        return jsonify({'error': 'Import batch exceeds maximum limit (2,000 records)'}), 400"""

assert p17_old in content, "Target 17 (confirm_fleet_calendar_import) pattern not found"
content = content.replace(p17_old, p17_new, 1)

with open(BACKEND_PATH, "w", encoding="utf-8") as f:
    f.write(content)

print(f"Successfully patched all 17 Round 6 backend targets! Original: {orig_len} chars, New: {len(content)} chars.")
