import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, PieChart, Reference
from datetime import datetime
import os

wb = openpyxl.Workbook()
ws_dash = wb.active
ws_dash.title = "Executive Summary"
ws_dash.views.sheetView[0].showGridLines = True

# Colors
PRIMARY_GREEN = "00B14F"
NAVY_HEADER = "1E293B"
CARD_BG = "F8FAFC"
TEXT_MUTED = "64748B"
BORDER_COLOR = "CBD5E1"

thin_border = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR)
)

report_title = "AUTORIDE SYSTEM - EXECUTIVE PERFORMANCE DASHBOARD"
date_subtitle = f"Period: All Time | Generated: {datetime.now().strftime('%b %d, %Y %I:%M %p')}"

# Row Heights
ws_dash.row_dimensions[1].height = 24
ws_dash.row_dimensions[2].height = 24
ws_dash.row_dimensions[3].height = 20
ws_dash.row_dimensions[4].height = 10
ws_dash.row_dimensions[5].height = 22
ws_dash.row_dimensions[6].height = 28
ws_dash.row_dimensions[7].height = 20
ws_dash.row_dimensions[8].height = 12
ws_dash.row_dimensions[25].height = 14

# Banner A1:L2
ws_dash.merge_cells("A1:L2")
ws_dash["A1"] = report_title
ws_dash["A1"].font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
ws_dash["A1"].fill = PatternFill(start_color=PRIMARY_GREEN, end_color=PRIMARY_GREEN, fill_type="solid")
ws_dash["A1"].alignment = Alignment(horizontal="center", vertical="center")

# Subtitle A3:L3
ws_dash.merge_cells("A3:L3")
ws_dash["A3"] = date_subtitle
ws_dash["A3"].font = Font(name="Calibri", size=10, italic=True, color=TEXT_MUTED)
ws_dash["A3"].alignment = Alignment(horizontal="center", vertical="center")

# 4 Horizontal KPI Cards
total_revenue = 1359397.0
total_bookings = 302
total_expenses = 125000.0
net_profit = total_revenue - total_expenses
avg_rev = total_revenue / total_bookings if total_bookings else 0
margin_pct = (net_profit / total_revenue * 100) if total_revenue > 0 else 0
active_fleet = 18

cards = [
    {
        "cols": (1, 3), # A:C
        "title": "TOTAL REVENUE",
        "value": f"PHP {total_revenue:,.2f}",
        "sub": f"Avg: PHP {avg_rev:,.2f} / booking",
        "val_color": PRIMARY_GREEN
    },
    {
        "cols": (4, 6), # D:F
        "title": "TOTAL BOOKINGS",
        "value": f"{total_bookings:,}",
        "sub": f"Fleet: {active_fleet} active vehicles",
        "val_color": "1E293B"
    },
    {
        "cols": (7, 9), # G:I
        "title": "TOTAL EXPENSES",
        "value": f"PHP {total_expenses:,.2f}",
        "sub": "Operations & Maintenance",
        "val_color": "DC2626"
    },
    {
        "cols": (10, 12), # J:L
        "title": "NET PROFIT",
        "value": f"PHP {net_profit:,.2f}",
        "sub": f"Operating Margin: {margin_pct:.1f}%",
        "val_color": PRIMARY_GREEN if net_profit >= 0 else "DC2626"
    }
]

header_fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
body_fill = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")

for card in cards:
    c_start, c_end = card["cols"]
    col_start_letter = get_column_letter(c_start)
    col_end_letter = get_column_letter(c_end)

    # Style all cells in card range
    for r in range(5, 8):
        for c in range(c_start, c_end + 1):
            cell = ws_dash.cell(row=r, column=c)
            cell.border = thin_border
            if r == 5:
                cell.fill = header_fill
            else:
                cell.fill = body_fill

    # Merge rows within the card
    ws_dash.merge_cells(f"{col_start_letter}5:{col_end_letter}5")
    ws_dash.merge_cells(f"{col_start_letter}6:{col_end_letter}6")
    ws_dash.merge_cells(f"{col_start_letter}7:{col_end_letter}7")

    top_cell = ws_dash[f"{col_start_letter}5"]
    top_cell.value = card["title"]
    top_cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    top_cell.alignment = Alignment(horizontal="center", vertical="center")

    val_cell = ws_dash[f"{col_start_letter}6"]
    val_cell.value = card["value"]
    val_cell.font = Font(name="Calibri", size=13, bold=True, color=card["val_color"])
    val_cell.alignment = Alignment(horizontal="center", vertical="center")

    sub_cell = ws_dash[f"{col_start_letter}7"]
    sub_cell.value = card["sub"]
    sub_cell.font = Font(name="Calibri", size=8, italic=True, color=TEXT_MUTED)
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

# Sheet 2: Chart Data
ws_data = wb.create_sheet(title="Chart_Data")
ws_data.views.sheetView[0].showGridLines = True
ws_data.append(["Period", "Bookings", "Revenue (PHP)"])

chart_trend = [
    ("2026-05", 25, 115000),
    ("2026-06", 42, 189000),
    ("2026-07", 58, 260000),
    ("2026-08", 65, 292000),
    ("2026-09", 72, 324000),
    ("2026-10", 40, 179397)
]
for row in chart_trend:
    ws_data.append(list(row))

for col in range(1, 4):
    cell = ws_data.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill

top_vehs = [
    {"brand": "Toyota", "model": "Fortuner", "plate": "ABC 1234", "rev": 345000},
    {"brand": "Mitsubishi", "model": "Montero", "plate": "XYZ 9876", "rev": 298000},
    {"brand": "Ford", "model": "Everest", "plate": "EVE 5544", "rev": 265000},
    {"brand": "Honda", "model": "Civic", "plate": "CIV 1122", "rev": 210000},
    {"brand": "Nissan", "model": "Navara", "plate": "NAV 3344", "rev": 150000}
]

ws_data.cell(row=1, column=5, value="Vehicle")
ws_data.cell(row=1, column=6, value="Revenue (PHP)")
ws_data.cell(row=1, column=5).font = Font(bold=True, color="FFFFFF")
ws_data.cell(row=1, column=6).font = Font(bold=True, color="FFFFFF")
ws_data.cell(row=1, column=5).fill = header_fill
ws_data.cell(row=1, column=6).fill = header_fill

for idx, tv in enumerate(top_vehs, start=2):
    ws_data.cell(row=idx, column=5, value=f"{tv['brand']} {tv['model']} ({tv['plate']})")
    ws_data.cell(row=idx, column=6, value=tv["rev"])

# Charts
# 1. Revenue Trend Line Chart (Left, A9)
line_chart = LineChart()
line_chart.title = "Revenue Trend"
line_chart.style = 13
line_chart.y_axis.title = "Revenue (PHP)"
line_chart.x_axis.title = "Period"
line_chart.width = 15.5
line_chart.height = 11

rev_data_ref = Reference(ws_data, min_col=3, min_row=1, max_row=len(chart_trend)+1)
cats_ref = Reference(ws_data, min_col=1, min_row=2, max_row=len(chart_trend)+1)
line_chart.add_data(rev_data_ref, titles_from_data=True)
line_chart.set_categories(cats_ref)
ws_dash.add_chart(line_chart, "A9")

# 2. Top Vehicles Pie Chart (Right, G9)
pie_chart = PieChart()
pie_chart.title = "Top Vehicles by Revenue"
pie_chart.width = 15.5
pie_chart.height = 11

pie_data = Reference(ws_data, min_col=6, min_row=1, max_row=len(top_vehs)+1)
pie_cats = Reference(ws_data, min_col=5, min_row=2, max_row=len(top_vehs)+1)
pie_chart.add_data(pie_data, titles_from_data=True)
pie_chart.set_categories(pie_cats)
ws_dash.add_chart(pie_chart, "G9")

# 3. Bookings Detail Bar Chart (Full Width, A26)
bar_chart = BarChart()
bar_chart.type = "col"
bar_chart.style = 10
bar_chart.title = "Bookings Detail"
bar_chart.y_axis.title = "Number of Bookings"
bar_chart.x_axis.title = "Period"
bar_chart.width = 31.5
bar_chart.height = 11

b_data_ref = Reference(ws_data, min_col=2, min_row=1, max_row=len(chart_trend)+1)
bar_chart.add_data(b_data_ref, titles_from_data=True)
bar_chart.set_categories(cats_ref)
ws_dash.add_chart(bar_chart, "A26")

# Set Page orientation and column widths
for ws in [ws_dash, ws_data]:
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    if ws == ws_dash:
        ws.page_setup.fitToHeight = 1
        for c_idx in range(1, 13):
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = 14
    else:
        ws.page_setup.fitToHeight = 0
        for col in ws.columns:
            max_len = 0
            for cell in col:
                v_str = str(cell.value or '')
                if len(v_str) > max_len and len(v_str) < 50:
                    max_len = len(v_str)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

out_path = os.path.join(os.path.dirname(__file__), "test_executive_dashboard.xlsx")
wb.save(out_path)
print(f"Successfully generated: {out_path}")
