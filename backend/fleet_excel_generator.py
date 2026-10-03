"""
Fleet Calendar Excel Export Generator
Generates an Excel workbook with:
1. 'Fleet Monitoring Calendar': A visual, color-coded, merged calendar grid matching ARC-MONITORING-2026.xlsx
2. 'Bookings Masterlist': A detailed tabular master list with full financial and customer details
"""
import io
import datetime
import calendar
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

# Default palette for bookings if no custom calendar_color is set
BOOKING_PALETTE = [
    '3D85C6', # Blue
    'FFD966', # Yellow
    '00FFFF', # Cyan
    'FF9900', # Orange
    'A4C2F4', # Soft Blue
    'B6D7A8', # Soft Green
    'D9D2E9', # Soft Purple
    'EA9999', # Soft Red
    'F9CB9C', # Soft Orange
    'FFE599', # Pale Yellow
    '9FC5E8', # Sky Blue
    '45818E', # Teal
    '6AA84F', # Green
    'B4A7D6', # Lilac
]

def clean_hex(hex_val, default='3D85C6'):
    if not hex_val:
        return default
    s = str(hex_val).strip().lstrip('#').upper()
    if len(s) == 6:
        return s
    if len(s) == 8: # ARGB
        return s[2:]
    return default

def generate_fleet_excel(vehicles, bookings, start_date=None, end_date=None, is_all=False):
    """
    vehicles: list of dicts with keys: id, brand, model, plate_number, status
    bookings: list of dicts with booking fields
    start_date: 'YYYY-MM-DD' or date object (optional)
    end_date: 'YYYY-MM-DD' or date object (optional)
    is_all: bool
    """
    wb = Workbook()
    
    # -------------------------------------------------------------
    # 1. Prepare Date Range and Months
    # -------------------------------------------------------------
    # Parse dates from bookings to find range if not specified
    parsed_bookings = []
    for b in bookings:
        b_copy = dict(b)
        s_date = b.get('start_date')
        e_date = b.get('end_date')
        if isinstance(s_date, str):
            try:
                s_date = datetime.datetime.strptime(s_date[:10], '%Y-%m-%d').date()
            except Exception:
                s_date = None
        elif isinstance(s_date, datetime.datetime):
            s_date = s_date.date()
            
        if isinstance(e_date, str):
            try:
                e_date = datetime.datetime.strptime(e_date[:10], '%Y-%m-%d').date()
            except Exception:
                e_date = None
        elif isinstance(e_date, datetime.datetime):
            e_date = e_date.date()
            
        b_copy['_start_date'] = s_date
        b_copy['_end_date'] = e_date or s_date
        parsed_bookings.append(b_copy)

    # Determine which (year, month) pairs to render
    months_to_render = [] # list of (year, month)
    
    if start_date and end_date and not is_all:
        if isinstance(start_date, str):
            start_date = datetime.datetime.strptime(start_date[:10], '%Y-%m-%d').date()
        if isinstance(end_date, str):
            end_date = datetime.datetime.strptime(end_date[:10], '%Y-%m-%d').date()
            
        curr_y = start_date.year
        curr_m = start_date.month
        end_y = end_date.year
        end_m = end_date.month
        while (curr_y < end_y) or (curr_y == end_y and curr_m <= end_m):
            months_to_render.append((curr_y, curr_m))
            if curr_m == 12:
                curr_y += 1
                curr_m = 1
            else:
                curr_m += 1
    else:
        # For "all" or unspecified, find min and max month from bookings, or default to current year
        years_months = set()
        for b in parsed_bookings:
            if b['_start_date']:
                years_months.add((b['_start_date'].year, b['_start_date'].month))
            if b['_end_date']:
                years_months.add((b['_end_date'].year, b['_end_date'].month))
                
        if years_months:
            min_ym = min(years_months)
            max_ym = max(years_months)
            curr_y, curr_m = min_ym
            end_y, end_m = max_ym
            while (curr_y < end_y) or (curr_y == end_y and curr_m <= end_m):
                months_to_render.append((curr_y, curr_m))
                if curr_m == 12:
                    curr_y += 1
                    curr_m = 1
                else:
                    curr_m += 1
        else:
            today = datetime.date.today()
            # Default to all 12 months of current year
            for m in range(1, 13):
                months_to_render.append((today.year, m))

    if not months_to_render:
        today = datetime.date.today()
        months_to_render = [(today.year, today.month)]

    # -------------------------------------------------------------
    # 2. Build Sheet 1: Fleet Monitoring Calendar
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Fleet Monitoring Calendar"
    ws1.views.sheetView[0].showGridLines = True
    
    # Styles matching ARC-MONITORING-2026.xlsx
    font_car_header = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
    fill_car_header = PatternFill(start_color="666666", end_color="666666", fill_type="solid")
    
    font_month_header = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
    fill_month_header = PatternFill(start_color="FF9900", end_color="FF9900", fill_type="solid")
    
    font_day_header = Font(name="Times New Roman", size=9, bold=True, color="000000")
    fill_day_header = PatternFill(start_color="B7B7B7", end_color="B7B7B7", fill_type="solid")
    
    font_vehicle = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
    fill_vehicle = PatternFill(start_color="6AA84F", end_color="6AA84F", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="CCCCCC")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    dark_border_side = Side(border_style="thin", color="555555")
    box_border = Border(left=dark_border_side, right=dark_border_side, top=dark_border_side, bottom=dark_border_side)

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_vehicle = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_booking = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Set column widths
    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 12
    ws1.column_dimensions['C'].width = 12
    for c_idx in range(4, 70): # D to BQ
        col_l = get_column_letter(c_idx)
        ws1.column_dimensions[col_l].width = 6.0

    current_row = 1

    # Map vehicles list
    # Sort vehicles nicely
    sorted_vehicles = sorted(vehicles, key=lambda v: (v.get('brand') or '', v.get('model') or '', v.get('plate_number') or ''))

    # Build blocks for each month
    for y, m in months_to_render:
        month_name = calendar.month_name[m].upper()
        num_days = calendar.monthrange(y, m)[1]
        
        # Max col for this month: Col 4 is Day 1 AM, Col 5 is Day 1 PM ...
        last_col = 3 + (num_days * 2) # e.g. for 31 days: 3 + 62 = 65 (BM)
        
        # Row 1: Header Row
        # A..C: CAR MAKE
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        c_car_make = ws1.cell(row=current_row, column=1, value="CAR MAKE")
        c_car_make.font = font_car_header
        c_car_make.fill = fill_car_header
        c_car_make.alignment = align_center
        
        # Style merged cells A1..C1
        for col in range(1, 4):
            c = ws1.cell(row=current_row, column=col)
            c.fill = fill_car_header
            c.border = box_border

        # D..last_col: Month Year Banner (e.g. "JANUARY 2026")
        ws1.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=last_col)
        c_month = ws1.cell(row=current_row, column=4, value=f"{month_name} {y}")
        c_month.font = font_month_header
        c_month.fill = fill_month_header
        c_month.alignment = align_center
        for col in range(4, last_col + 1):
            c = ws1.cell(row=current_row, column=col)
            c.fill = fill_month_header
            c.border = box_border
            
        ws1.row_dimensions[current_row].height = 24
        current_row += 1

        # Row 2: Day of Week Names (e.g. MON, TUE, WED)
        # Row 3: Day Numbers (1, 2, 3...)
        row_day_name = current_row
        row_day_num = current_row + 1
        
        # Col A..C in day rows
        ws1.merge_cells(start_row=row_day_name, start_column=1, end_row=row_day_num, end_column=3)
        c_corner = ws1.cell(row=row_day_name, column=1, value="")
        c_corner.fill = PatternFill(start_color="EFEFEF", end_color="EFEFEF", fill_type="solid")
        for r_sub in (row_day_name, row_day_num):
            for col in range(1, 4):
                ws1.cell(row=r_sub, column=col).border = box_border

        day_abbrs = ['MON', 'TUE', 'WED', 'THU', 'FRI', 'SAT', 'SUN']
        for d in range(1, num_days + 1):
            day_dt = datetime.date(y, m, d)
            day_name = day_abbrs[day_dt.weekday()]
            
            c_start = 4 + (d - 1) * 2
            c_end = c_start + 1
            
            # Merge 2 columns for the day
            ws1.merge_cells(start_row=row_day_name, start_column=c_start, end_row=row_day_name, end_column=c_end)
            c_dn = ws1.cell(row=row_day_name, column=c_start, value=day_name)
            c_dn.font = font_day_header
            c_dn.fill = fill_day_header
            c_dn.alignment = align_center
            ws1.cell(row=row_day_name, column=c_end).fill = fill_day_header
            ws1.cell(row=row_day_name, column=c_start).border = box_border
            ws1.cell(row=row_day_name, column=c_end).border = box_border

            ws1.merge_cells(start_row=row_day_num, start_column=c_start, end_row=row_day_num, end_column=c_end)
            c_dnum = ws1.cell(row=row_day_num, column=c_start, value=d)
            c_dnum.font = font_day_header
            c_dnum.fill = fill_day_header
            c_dnum.alignment = align_center
            ws1.cell(row=row_day_num, column=c_end).fill = fill_day_header
            ws1.cell(row=row_day_num, column=c_start).border = box_border
            ws1.cell(row=row_day_num, column=c_end).border = box_border

        ws1.row_dimensions[row_day_name].height = 18
        ws1.row_dimensions[row_day_num].height = 18
        current_row += 2

        # Data Rows for each vehicle
        color_idx = 0
        for veh in sorted_vehicles:
            veh_id = veh.get('id')
            v_brand = veh.get('brand') or ''
            v_model = veh.get('model') or ''
            v_plate = veh.get('plate_number') or ''
            
            # Format vehicle label e.g.: VIOS J MT - NCX 4117
            veh_title = f"{v_brand} {v_model} - {v_plate}".strip().lstrip('-').strip()
            if not veh_title:
                veh_title = f"Vehicle #{veh_id}"

            # Left column A..C
            ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
            c_veh = ws1.cell(row=current_row, column=1, value=veh_title)
            c_veh.font = font_vehicle
            c_veh.fill = fill_vehicle
            c_veh.alignment = align_vehicle
            for col in range(1, 4):
                ws1.cell(row=current_row, column=col).fill = fill_vehicle
                ws1.cell(row=current_row, column=col).border = box_border

            # Set thin borders for all day cells in this row
            for col in range(4, last_col + 1):
                c = ws1.cell(row=current_row, column=col)
                c.border = thin_border

            # Find bookings for this vehicle that overlap this month
            veh_bookings = []
            for b in parsed_bookings:
                if str(b.get('vehicle_id')) == str(veh_id):
                    bs = b.get('_start_date')
                    be = b.get('_end_date')
                    if bs and be:
                        # Check overlap with month (y, m, 1) .. (y, m, num_days)
                        month_start = datetime.date(y, m, 1)
                        month_end = datetime.date(y, m, num_days)
                        if bs <= month_end and be >= month_start:
                            veh_bookings.append(b)

            # Sort vehicle bookings by start date
            veh_bookings.sort(key=lambda x: x['_start_date'])

            # Track occupied columns on this row to prevent openpyxl MergedCell conflicts
            occupied_cols = set()

            # Plot each booking across columns
            for bk in veh_bookings:
                bs = bk['_start_date']
                be = bk['_end_date']
                
                # Clamp within this month
                s_day = max(1, bs.day if (bs.year == y and bs.month == m) else 1)
                e_day = min(num_days, be.day if (be.year == y and be.month == m) else num_days)
                
                col_start = 4 + (s_day - 1) * 2
                col_end = 4 + (e_day - 1) * 2 + 1
                
                # If col_start is occupied, shift forward to next free column
                while col_start <= col_end and col_start in occupied_cols:
                    col_start += 1
                
                if col_start > col_end or col_start > last_col:
                    continue
                    
                # Clamp col_end before any occupied column
                for c_check in range(col_start, col_end + 1):
                    if c_check in occupied_cols:
                        col_end = c_check - 1
                        break
                        
                col_end = min(col_end, last_col)
                if col_end < col_start:
                    continue

                # Pick booking color
                bk_color = bk.get('calendar_color')
                if bk_color:
                    hex_color = clean_hex(bk_color)
                else:
                    hex_color = BOOKING_PALETTE[color_idx % len(BOOKING_PALETTE)]
                    color_idx += 1
                    
                bk_fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
                
                # Determine font color for readability (black or white)
                r = int(hex_color[0:2], 16)
                g = int(hex_color[2:4], 16)
                b_lum = int(hex_color[4:6], 16)
                lum = (0.299 * r + 0.587 * g + 0.114 * b_lum)
                font_c = "000000" if lum > 140 else "FFFFFF"
                bk_font = Font(name="Times New Roman", size=9, bold=True, color=font_c)

                # Customer name or label
                cust_name = str(bk.get('customer_name') or 'BOOKED').strip().upper()
                bk_text = cust_name

                # 1. Set value and style on the primary top-left cell FIRST
                top_left = ws1.cell(row=current_row, column=col_start, value=bk_text)
                top_left.font = bk_font
                top_left.fill = bk_fill
                top_left.alignment = align_booking
                top_left.border = box_border

                # Attach hover comment note with full booking details (matching ARC Monitoring)
                c_name = str(bk.get('customer_name') or 'Customer').strip()
                c_phone = str(bk.get('customer_phone') or 'N/A').strip()
                c_addr = str(bk.get('customer_address') or 'N/A').strip()
                c_em_name = str(bk.get('emergency_contact_name') or 'N/A').strip()
                c_em_phone = str(bk.get('emergency_contact_phone') or 'N/A').strip()
                c_unit = veh_title

                try:
                    c_pick_date = bs.strftime('%B %d, %Y')
                except Exception:
                    c_pick_date = str(bk.get('start_date') or '')

                try:
                    c_ret_date = be.strftime('%B %d, %Y')
                except Exception:
                    c_ret_date = str(bk.get('end_date') or '')

                c_pick_time = str(bk.get('start_time') or '06:00 AM')
                c_ret_time = str(bk.get('end_time') or '06:00 AM')
                c_dest = str(bk.get('destination') or 'N/A')
                c_purpose = str(bk.get('rental_purpose') or 'Travel / Rental')
                c_total = f"PHP {float(bk.get('total_price') or 0):,.2f}"
                c_status = str(bk.get('status') or 'Active').capitalize()

                comment_lines = [
                    f"Name: {c_name}",
                    f"Contact no: {c_phone}",
                    f"Complete Address: {c_addr}",
                    f"Emergency contact person: {c_em_name}",
                    f"Contact no: {c_em_phone}",
                    f"Unit needed: {c_unit}",
                    f"Date of pick up: {c_pick_date}",
                    f"Date of return: {c_ret_date}",
                    f"Time of pick up: {c_pick_time}",
                    f"Time of return: {c_ret_time}",
                    f"Destination: {c_dest}",
                    f"Purpose of Rental: {c_purpose}",
                    f"Total Amount: {c_total} ({c_status})"
                ]
                cell_comment = Comment("\n".join(comment_lines), "Autoride")
                cell_comment.width = 240
                cell_comment.height = 190
                top_left.comment = cell_comment

                # 2. Merge if spanning more than 1 cell
                if col_end > col_start:
                    ws1.merge_cells(start_row=current_row, start_column=col_start, end_row=current_row, end_column=col_end)
                
                # 3. Apply border & fill to the remaining cells in the span and mark occupied
                for c_span in range(col_start, col_end + 1):
                    occupied_cols.add(c_span)
                    cell_s = ws1.cell(row=current_row, column=c_span)
                    cell_s.fill = bk_fill
                    cell_s.border = box_border

            ws1.row_dimensions[current_row].height = 24
            current_row += 1

        # Leave 2 blank rows after each month block
        current_row += 2

    # -------------------------------------------------------------
    # 3. Build Sheet 2: Bookings Masterlist (Detailed Tabular View)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Bookings Masterlist")
    ws2.views.sheetView[0].showGridLines = True

    headers = [
        "Booking ID", "Status", "Vehicle / Unit", "Plate Number", 
        "Customer Name", "Contact Phone", "Customer Email", "Address",
        "Start Date", "Start Time", "End Date", "End Time",
        "Rental Type", "Base Price", "Driver Fee", "Payment Fee",
        "Total Price", "Amount Paid", "Balance Amount", "Payment Status",
        "Payment Method", "Destination", "Emergency Contact", "Emergency Phone"
    ]

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws2.row_dimensions[1].height = 26
    for col_i, h in enumerate(headers, start=1):
        c = ws2.cell(row=1, column=col_i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = box_border

    data_font = Font(name="Segoe UI", size=9)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_cen = Alignment(horizontal="center", vertical="center")

    for row_idx, bk in enumerate(parsed_bookings, start=2):
        ws2.row_dimensions[row_idx].height = 20
        v_title = f"{bk.get('brand') or ''} {bk.get('model') or ''}".strip()
        row_values = [
            f"#{bk.get('id')}",
            str(bk.get('status') or '').capitalize(),
            v_title or f"Vehicle #{bk.get('vehicle_id')}",
            bk.get('plate_number') or '',
            bk.get('customer_name') or '',
            bk.get('customer_phone') or '',
            bk.get('customer_email') or '',
            bk.get('customer_address') or '',
            str(bk.get('start_date') or ''),
            str(bk.get('start_time') or ''),
            str(bk.get('end_date') or ''),
            str(bk.get('end_time') or ''),
            str(bk.get('rental_type') or '').capitalize(),
            float(bk.get('base_price') or 0),
            float(bk.get('driver_fee') or 0),
            float(bk.get('payment_fee') or 0),
            float(bk.get('total_price') or 0),
            float(bk.get('amount_paid') or 0),
            float(bk.get('balance_amount') or 0),
            str(bk.get('payment_status') or '').capitalize(),
            bk.get('payment_method') or 'Cash',
            bk.get('destination') or '',
            bk.get('emergency_contact_name') or '',
            bk.get('emergency_contact_phone') or '',
        ]

        for col_i, val in enumerate(row_values, start=1):
            c = ws2.cell(row=row_idx, column=col_i, value=val)
            c.font = data_font
            c.border = thin_border
            if col_i in (1, 2, 9, 10, 11, 12, 13, 20):
                c.alignment = align_cen
            elif col_i in (14, 15, 16, 17, 18, 19):
                c.alignment = align_right
                c.number_format = '#,##0.00'
            else:
                c.alignment = align_left

    # Auto fit column widths for Sheet 2
    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 11)

    # Return as BytesIO
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
